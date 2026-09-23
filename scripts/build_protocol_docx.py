from __future__ import annotations

import argparse
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "ESP32-S3_USB_Modbus_IO_Protocol.docx"
BOARD_IMAGE = ROOT / "docs" / "assets" / "esp32-s3-board.png"

NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
CODE_FILL = "F3F5F7"
BLACK = RGBColor(0, 0, 0)


DIGITAL_CHANNELS = [
    (0, 0, "启动绑带和 BOOT 按键"),
    (1, 1, "ADC 模拟通道 0"),
    (2, 2, "ADC 模拟通道 1"),
    (3, 3, "ADC 模拟通道 2，启动绑带"),
    (4, 4, "ADC 模拟通道 3"),
    (5, 5, "ADC 模拟通道 4"),
    (6, 6, "ADC 模拟通道 5"),
    (7, 7, "ADC 模拟通道 6"),
    (8, 8, "ADC 模拟通道 7"),
    (9, 9, "ADC 模拟通道 8"),
    (10, 10, "ADC 模拟通道 9"),
    (11, 11, "ADC 模拟通道 10"),
    (12, 12, "ADC 模拟通道 11"),
    (13, 13, "ADC 模拟通道 12"),
    (14, 14, "ADC 模拟通道 13"),
    (15, 15, "ADC 模拟通道 14"),
    (16, 16, "ADC 模拟通道 15"),
    (17, 17, "ADC 模拟通道 16"),
    (18, 18, "ADC 模拟通道 17"),
    (19, 21, "通用 IO"),
    (20, 38, "通用 IO"),
    (21, 39, "JTAG 复用"),
    (22, 40, "JTAG 复用"),
    (23, 41, "JTAG 复用"),
    (24, 42, "JTAG 复用"),
    (25, 43, "板上标记 TX，UART0 复用"),
    (26, 44, "板上标记 RX，UART0 复用"),
    (27, 45, "启动绑带"),
    (28, 46, "启动绑带"),
    (29, 47, "通用 IO"),
    (30, 48, "通常连接板载 RGB LED"),
    (31, 35, "默认禁用，可能连接八线 PSRAM"),
    (32, 36, "默认禁用，可能连接八线 PSRAM"),
    (33, 37, "默认禁用，可能连接八线 PSRAM"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def style_run(run, *, font="Microsoft YaHei", size=10.5, bold=False, color=BLACK) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Microsoft YaHei")
    fonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    r_pr.extend((fonts, color, underline, size))
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), LIGHT_GRAY)


def add_table(doc, headers, rows, widths, alignments=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    set_cant_split(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        style_run(paragraph.add_run(str(label)), size=9.3, bold=True, color=RGBColor(255, 255, 255))

    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        for column, value in enumerate(values):
            cell = row.cells[column]
            set_cell_width(cell, widths[column])
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            if alignments:
                paragraph.alignment = alignments[column]
            style_run(paragraph.add_run(str(value)), size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_body(doc, text: str, *, bold_lead: str | None = None):
    paragraph = doc.add_paragraph(style="Body Text")
    paragraph.paragraph_format.keep_together = False
    if bold_lead:
        paragraph.paragraph_format.keep_with_next = True
        style_run(paragraph.add_run(bold_lead), bold=True)
    style_run(paragraph.add_run(text))
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        style_run(paragraph.add_run(item))


def add_code(doc, lines):
    for line in lines:
        paragraph = doc.add_paragraph(style="Code")
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        p_pr = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), CODE_FILL)
        p_pr.append(shading)
        style_run(paragraph.add_run(line), font="Cascadia Mono", size=9.2)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    style_run(run, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.32)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    body = styles["Body Text"]
    body.base_style = normal
    body.paragraph_format.first_line_indent = Cm(0.74)
    body.paragraph_format.space_after = Pt(6)
    body.paragraph_format.line_spacing = 1.2

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Microsoft YaHei"
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    subtitle.font.size = Pt(12)
    subtitle.font.color.rgb = BLACK
    subtitle.paragraph_format.space_after = Pt(10)

    for style_name, size, before, after in (
        ("Heading 1", 16, 14, 7),
        ("Heading 2", 12.5, 10, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    list_bullet = styles["List Bullet"]
    list_bullet.base_style = normal
    list_bullet.paragraph_format.left_indent = Cm(0.74)
    list_bullet.paragraph_format.first_line_indent = Cm(-0.42)
    list_bullet.paragraph_format.space_after = Pt(3)

    code = styles.add_style("Code", 1)
    code.font.name = "Cascadia Mono"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code.font.size = Pt(9.2)
    code.font.color.rgb = BLACK
    code.paragraph_format.left_indent = Cm(0.35)
    code.paragraph_format.right_indent = Cm(0.35)
    code.paragraph_format.line_spacing = 1.05

    footer_table = section.footer.add_table(rows=1, cols=2, width=Inches(7.06))
    footer_table.autofit = False
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = footer_table.rows[0].cells
    set_cell_width(left, 5.8)
    set_cell_width(right, 1.26)
    style_run(left.paragraphs[0].add_run("ESP32-S3 USB Modbus IO 协议 1.2"), size=8.5)
    add_page_number(right.paragraphs[0])


def preserve_board_image(board_image: Path) -> Path:
    """Migrate the original embedded photo once; retain a portable source asset."""
    if board_image.exists():
        return board_image
    if board_image == BOARD_IMAGE and OUTPUT.exists():
        with ZipFile(OUTPUT) as existing:
            photo = existing.read("word/media/image1.png")
        board_image.parent.mkdir(parents=True, exist_ok=True)
        board_image.write_bytes(photo)
        return board_image
    raise FileNotFoundError(f"Board image missing: {board_image}; use --board-image PATH")


def build(board_image: Path = BOARD_IMAGE) -> None:
    board_image = preserve_board_image(board_image)
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "ESP32-S3 USB Modbus IO 通信协议"
    doc.core_properties.subject = "ESP-IDF 固件通信与寄存器说明"
    doc.core_properties.author = "ESP32-S3 USB Modbus IO Project"
    doc.core_properties.keywords = "ESP32-S3, ESP-IDF, Modbus RTU, USB CDC, GPIO, ADC, PWM, I2C, SPI, UART"

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(title.add_run("ESP32-S3 USB Modbus IO 通信协议"), size=24, bold=True)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(subtitle.add_run("ESP-IDF 固件通信与寄存器说明"), size=12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(meta.add_run("协议版本 1.2    固件版本 1.2    从站地址 1"), size=10, bold=True)

    add_body(
        doc,
        "本手册给出图示 ESP32-S3-WROOM-1 开发板的 USB 虚拟串口 Modbus RTU 接口。"
        "主站可读取数字输入和输出状态、控制数字输出、读取 GPIO1 到 GPIO18 的 ADC 数据，"
        "并配置 PWM、I²C、SPI 和 UART。CDC0 承载 Modbus 控制事务，CDC1 提供 UART 原始字节透传。"
        "版本 1.2 保留 1.0 和 1.1 的数字 IO、ADC 和 PWM 地址。",
    )

    if board_image.exists():
        image_paragraph = doc.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture = image_paragraph.add_run().add_picture(str(board_image), width=Inches(5.35))
        picture._inline.docPr.set("descr", "ESP32-S3-WROOM-1 双 USB-C 开发板引脚图")
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(caption.add_run("图 1  ESP32-S3-WROOM-1 开发板引脚"), size=9)

    doc.add_page_break()

    doc.add_heading("1 快速使用", level=1)
    add_body(
        doc,
        "通信线应连接到 GPIO19 和 GPIO20 对应的原生 USB 或 OTG 接口。另一 USB-C 接口如果经过"
        "USB 转串口芯片，只能用于 UART 烧录或调试，不承载本协议。USB CDC 的波特率设置不影响"
        "实际 USB 传输，但主站仍需发送完整的 Modbus RTU 帧和 CRC。",
    )
    add_table(
        doc,
        ["项目", "默认值", "说明"],
        [
            ["目标芯片", "ESP32-S3", "ESP32-S3-WROOM-1 开发板"],
            ["传输", "双 USB CDC-ACM", "CDC0 Modbus 控制；CDC1 UART 字节透传"],
            ["协议", "Modbus RTU", "CRC16 低字节先发送"],
            ["从站地址", "1", "可在 menuconfig 中改为 1 到 247"],
            ["不完整帧超时", "20 ms", "完整帧立即处理"],
            ["数字通道", "34", "默认可用 31 个，GPIO35 到 GPIO37 保留"],
            ["模拟通道", "18", "GPIO1 到 GPIO18"],
            ["PWM 输出", "最多 8 路", "同时最多 4 种频率；10 Hz 到 100 kHz"],
            ["调试总线", "I²C SPI UART", "经功能码 0x41 配置，通道引脚可选择"],
        ],
        [1.45, 1.55, 4.0],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_code(
        doc,
        [
            "idf.py set-target esp32s3",
            "idf.py build",
            "idf.py -p COMx flash",
            "python -m pip install -r requirements.txt",
            "python tools/modbus_usb_client.py --port COM8 info",
        ],
    )

    doc.add_heading("2 硬件和电气要求", level=1)
    add_bullets(
        doc,
        [
            "GPIO 逻辑电平为 3.3 V，不得直接输入 5 V。",
            "GPIO1 到 GPIO18 支持 ADC。建议使用低阻抗信号源，并保证输入始终处于芯片允许范围。",
            "继电器、线圈、电机、灯带和其他大电流负载必须通过驱动器和保护电路连接。",
            "GPIO0、GPIO3、GPIO45、GPIO46 是启动绑带引脚，外部电路不能在复位期间强迫错误电平。",
            "GPIO0 通常连接 BOOT 按键，GPIO48 通常连接板载 RGB LED。",
            "GPIO35 到 GPIO37 可能连接八线 PSRAM，默认禁用。只有确认模组接线兼容后才能启用。",
            "I²C 需要外部上拉；上拉不得接至 5 V。SPI 和 UART 跨电压域使用电平转换。",
        ],
    )

    heading = doc.add_heading("3 RTU 帧和 USB 传输", level=1)
    heading.paragraph_format.page_break_before = True
    add_table(
        doc,
        ["字段", "长度", "说明"],
        [
            ["从站地址", "1 字节", "默认 0x01；地址 0 仅用于广播写"],
            ["功能码", "1 字节", "见第 4 节"],
            ["数据", "可变", "多字节数值高字节在前"],
            ["CRC16", "2 字节", "多项式 0xA001，低字节先发送"],
        ],
        [1.35, 1.1, 4.55],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body(
        doc,
        "USB 可能拆分或合并串口写操作。固件按功能码和长度字段重组请求。完整帧立即处理；"
        "不完整字节流超过配置超时后丢弃。主站应在一次写操作中提交完整 RTU 帧，并等待响应后"
        "再发送下一条请求。地址 0 的合法写请求会执行但不响应，广播读会被忽略。",
    )

    doc.add_heading("4 功能码", level=1)
    add_table(
        doc,
        ["功能码", "名称", "用途"],
        [
            ["0x01", "Read Coils", "读取数字输出命令状态"],
            ["0x02", "Read Discrete Inputs", "读取 GPIO 实际逻辑电平"],
            ["0x03", "Read Holding Registers", "读取模式、GPIO 映射、设备信息和 PWM 配置"],
            ["0x04", "Read Input Registers", "读取 ADC 原始值或校准毫伏值"],
            ["0x05", "Write Single Coil", "设置一个数字输出并自动进入输出模式"],
            ["0x06", "Write Single Register", "设置一个 GPIO 模式"],
            ["0x0F", "Write Multiple Coils", "批量设置数字输出"],
            ["0x10", "Write Multiple Registers", "批量设置模式或写单路完整 PWM 配置"],
            ["0x41", "Debugger Extension", "I²C SPI UART 配置及事务，仅单播"],
        ],
        [1.0, 2.45, 3.55],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_heading("5 数字通道映射", level=1)
    add_body(
        doc,
        "Modbus PDU 地址从 0 开始，传统线圈号和离散输入号从 1 开始显示。线圈返回最近写入的"
        "输出命令状态，离散输入返回引脚实际电平。GPIO35 到 GPIO37 的通道号始终保留，但默认"
        "访问会返回非法数据地址异常。PWM 模式读线圈仍返回最近的数字输出命令值；读离散输入"
        "只得到采样瞬间电平，不能用来判断 PWM 占空比。模拟或总线占用模式读离散输入返回异常 0x04。",
    )
    digital_rows = []
    for channel, gpio, note in DIGITAL_CHANNELS:
        default = "禁用" if channel >= 31 else "浮空输入"
        digital_rows.append(
            [channel, gpio, f"{channel + 1:05d}", f"{10001 + channel}", default, note]
        )
    add_table(
        doc,
        ["通道", "GPIO", "线圈号", "离散输入号", "默认", "备注"],
        digital_rows,
        [0.55, 0.6, 0.85, 1.05, 0.9, 3.05],
        [
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
        font_size=8.7,
    )

    doc.add_heading("6 模拟输入寄存器", level=1)
    add_body(
        doc,
        "模拟通道 0 到 17 依次对应 GPIO1 到 GPIO18。ADC 使用 12 dB 衰减和默认位宽。读取时会"
        "把非输出且未被总线占用的 GPIO 切换为模拟模式。如果对应 GPIO 正在数字输出、PWM 输出或被外设占用，固件返回异常"
        " 0x04，保持现有输出；可先明确切换到模拟模式再读取 ADC。",
    )
    add_table(
        doc,
        ["PDU 地址", "传统寄存器号", "内容", "数值"],
        [
            ["0x0000 到 0x0011", "30001 到 30018", "ADC 原始值", "典型范围 0 到 4095"],
            ["0x0100 到 0x0111", "30257 到 30274", "校准电压", "mV；无校准时 0xFFFF"],
        ],
        [1.65, 1.55, 1.4, 2.4],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
         WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body(
        doc,
        "校准毫伏值来自 ESP-IDF ADC 曲线拟合校准接口。没有校准数据时，原始寄存器仍可读取，"
        "毫伏寄存器返回 0xFFFF。ADC 结果会受信号源阻抗、噪声、板级布局和芯片误差影响。",
    )

    heading = doc.add_heading("7 保持寄存器", level=1)
    heading.paragraph_format.page_break_before = True
    doc.add_heading("7 1 GPIO 模式", level=2)
    add_body(
        doc,
        "PDU 地址 0x0000 到 0x0021 对应数字通道 0 到 33。可用 0x03 读取，用 0x06 或 0x10 写入。"
        "写线圈会停止 PWM 并切换到数字输出模式；PWM 通道切换为模式 0 到 4 时释放 PWM 资源。"
        "切换到模式 3 时恢复最近一次线圈命令值。",
    )
    add_table(
        doc,
        ["模式值", "含义", "限制"],
        [
            [0, "浮空数字输入", "所有可用通道"],
            [1, "上拉数字输入", "所有可用通道"],
            [2, "下拉数字输入", "所有可用通道"],
            [3, "数字输出", "保留最近一次线圈命令值"],
            [4, "模拟输入", "仅数字通道 1 到 18"],
            [5, "PWM 输出", "使用已保存配置；初始 1000 Hz、50%"],
            [6, "I²C 占用", "只读，由 0x41 配置和释放"],
            [7, "SPI 占用", "只读，由 0x41 配置和释放"],
            [8, "UART 占用", "只读，由 0x41 配置和释放"],
        ],
        [1.0, 2.2, 3.8],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    doc.add_heading("7 2 GPIO 映射", level=2)
    add_body(
        doc,
        "PDU 地址 0x0100 到 0x0121 只读，对应数字通道 0 到 33。寄存器值是实际 GPIO 编号。"
        "默认禁用的 GPIO35 到 GPIO37 返回非法数据地址异常。",
    )
    doc.add_heading("7 3 设备信息", level=2)
    add_table(
        doc,
        ["PDU 地址", "传统寄存器号", "内容", "当前值"],
        [
            ["0x0200", "40513", "协议版本", "0x0102 表示 1.2"],
            ["0x0201", "40514", "固件版本", "0x0102 表示 1.2"],
            ["0x0202", "40515", "数字通道总数", "34"],
            ["0x0203", "40516", "模拟通道总数", "18"],
            ["0x0204", "40517", "能力位", "bit0 ADC 校准；bit1 GPIO35～37；bit2 PWM；bit3 I²C；bit4 SPI；bit5 UART"],
            ["0x0205", "40518", "最大 PWM 输出数", "8"],
            ["0x0206", "40519", "最大频率种类", "4"],
        ],
        [1.0, 1.2, 1.8, 3.0],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
         WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_body(doc, "前 5 个信息寄存器布局保持兼容。先读 0x0200 到 0x0204，确认 bit2 后再读取 PWM 扩展；确认 bit3 到 bit5 后再使用功能码 0x41。旧版 1.0 固件没有 PWM 扩展地址。")

    doc.add_heading("7 4 PWM 配置", level=2)
    add_body(doc, "数字通道 N 的 PWM 记录首地址为 0x0300 + 4 × N，每条记录占 4 个寄存器。整个配置区为 0x0300 到 0x0387，传统寄存器号为 40769 到 40904。通道 31 到 33 仍受 GPIO35 到 GPIO37 构建开关限制，禁用时返回 0x02。")
    add_table(
        doc,
        ["偏移", "内容", "范围和含义"],
        [
            ["+0", "频率高 16 位", "与低字拼为 32 位无符号整数，单位 Hz"],
            ["+1", "频率低 16 位", "合成频率为 10 到 100000 Hz"],
            ["+2", "占空比", "0 到 10000；每单位 0.01%；5000 表示 50%"],
            ["+3", "使能", "0 停止，1 启动；回读表示当前 PWM 状态"],
        ],
        [0.7, 1.6, 4.7],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_body(doc, "频率等于高字 × 65536 + 低字。1000 Hz 编码为 0x0000 0x03E8，100000 Hz 编码为 0x0001 0x86A0。每个寄存器内部高字节在前。读功能码 0x03 可读取单个字段或连续字段。")
    add_body(doc, "写功能码 0x10 必须从记录首地址开始，数量为 4、字节数为 8，一次提交频率、占空比和使能。禁止部分写或跨记录写；0x06 写 PWM 配置返回 0x03。参数超范围、使能不是 0 或 1 时，也返回 0x03。")
    add_body(doc, "使能为 1 时应用参数并进入模式 5。使能为 0 时保存参数；若通道原来为 PWM，则停止并切换为浮空输入，其他模式保持原样。停止后的参数可再次通过模式 5 启动。参数只保存在内存中，重启后 GPIO 恢复浮空输入，PWM 全部关闭，各通道 PWM 配置恢复为 1000 Hz、50%。")

    doc.add_heading("7 5 PWM 资源和精度", level=2)
    add_body(doc, "LEDC 硬件最多同时输出 8 路、使用 4 种频率。同频通道共享定时器，占空比独立。修改其中一路的占空比不会修改其他通道；不同频率需要另一只定时器。资源不足返回 0x06，资源检查阶段保持原有输出。可使用已有频率，或先停止不再使用的输出以释放资源。")
    add_body(doc, "时钟源为 40 MHz XTAL，驱动按频率选择 8 到 14 位计数分辨率。0.01% 是协议输入步进，实际占空比会量化到硬件刻度；频率越高，可用分辨率越低。实际频率也有分频量化误差。回读返回请求参数，不是波形测量值。0% 常低、100% 常高，两者仍占用 PWM 资源。")
    add_body(doc, "PWM 是 3.3 V 数字脉冲，不是 DAC 模拟电压。外部电路需要平滑电压时应按负载设计滤波和缓冲；需要精确频率、占空比或切换时序时，应在目标板上用示波器核验。")

    doc.add_heading("8 请求与响应规则", level=1)
    add_body(doc, "功能码 0x01 和 0x02 的位从响应数据字节最低位开始排列，未使用的最高位填 0。")
    add_body(doc, "功能码 0x05 使用 0xFF00 表示高电平，使用 0x0000 表示低电平。其他值返回异常 0x03。")
    add_body(doc, "功能码 0x03 和 0x04 的寄存器按高字节在前排列。0x06 只允许写 GPIO 模式地址，包括模式 5；写 PWM 配置返回 0x03。")
    add_body(doc, "0x0F 线圈值按最低位优先排列；0x10 每个寄存器占 2 字节。批量写先验证完整请求，成功响应回显地址和数量。")
    add_body(doc, "GPIO 模式区批量写若包含模式 5，数量必须为 1；否则整帧返回 0x03，不会部分启动。模式 0～4 可批量写。PWM 配置区只接受一条对齐的完整 4 寄存器记录。")

    doc.add_heading("9 异常响应", level=1)
    exception_table = add_table(
        doc,
        ["异常码", "名称", "含义"],
        [
            ["0x01", "Illegal Function", "不支持该功能码"],
            ["0x02", "Illegal Data Address", "地址越界、通道禁用或能力不匹配"],
            ["0x03", "Illegal Data Value", "数量、字节数、线圈、模式或 PWM 参数和格式无效"],
            ["0x04", "Server Device Failure", "模式冲突，或 GPIO ADC PWM I²C SPI UART 驱动失败"],
            ["0x06", "Server Device Busy", "PWM 通道或定时器资源不足"],
        ],
        [1.0, 2.1, 3.9],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    for row in exception_table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=35, bottom=35)
    doc.add_heading("10 报文示例", level=1)
    add_body(doc, "以下十六进制字节均包含 CRC，CRC 低字节在前。")
    examples = [
        ("读取 5 个设备信息寄存器", ["请求  01 03 02 00 00 05 84 71"]),
        ("读取数字通道 0 到 7 的实际电平", ["请求  01 02 00 00 00 08 79 CC", "响应  01 02 01 XX CRC_LO CRC_HI"]),
        ("把数字通道 1 设置为输出模式", ["请求  01 06 00 01 00 03 98 0B", "响应  01 06 00 01 00 03 98 0B"]),
        ("把数字通道 1 对应的 GPIO1 输出高电平", ["请求  01 05 00 01 FF 00 DD FA", "响应  01 05 00 01 FF 00 DD FA"]),
        ("读取模拟通道 0 的原始值", ["请求  01 04 00 00 00 01 31 CA", "响应  01 04 02 RAW_H RAW_L CRC_LO CRC_HI"]),
        ("读取模拟通道 0 的校准毫伏值", ["请求  01 04 01 00 00 01 30 36", "响应  01 04 02 MV_H MV_L CRC_LO CRC_HI"]),
        ("通道 1 GPIO1 设置 1000 Hz、50% 并启动 PWM", ["请求  01 10 03 04 00 04 08 00 00 03 E8 13 88 00 01 67 3C", "响应  01 10 03 04 00 04 80 4F"]),
        ("读取通道 1 PWM 配置，响应对应上例启动后的状态", ["请求  01 03 03 04 00 04 05 8C", "响应  01 03 08 00 00 03 E8 13 88 00 01 B0 9D"]),
        ("停止通道 1 PWM，恢复浮空输入并保留参数", ["请求  01 10 03 04 00 04 08 00 00 03 E8 13 88 00 00 A6 FC", "响应  01 10 03 04 00 04 80 4F"]),
    ]
    for label, lines in examples:
        add_body(doc, "", bold_lead=label)
        add_code(doc, lines)

    doc.add_heading("11 主机端测试", level=1)
    add_body(doc, "仓库中的 tools/modbus_usb_client.py 提供常用命令。安装 pyserial 后指定 CDC 串口。")
    add_code(
        doc,
        [
            "python -m pip install -r requirements.txt",
            "python tools/modbus_usb_client.py --port COM8 info",
            "python tools/modbus_usb_client.py --port COM8 read-map 0 31",
            "python tools/modbus_usb_client.py --port COM8 read-inputs 0 8",
            "python tools/modbus_usb_client.py --port COM8 write-output 1 on",
            "python tools/modbus_usb_client.py --port COM8 read-analog 0 4",
            "python tools/modbus_usb_client.py --port COM8 read-analog 0 4 --raw",
            "python tools/modbus_usb_client.py --port COM8 set-mode 1 0",
            "python tools/modbus_usb_client.py --port COM8 pwm-set 1 1000 50",
            "python tools/modbus_usb_client.py --port COM8 pwm-read 1",
            "python tools/modbus_usb_client.py --port COM8 pwm-stop 1",
            "python tools/modbus_usb_client.py --port COM8 pwm-set 1 1000 25 --disabled",
            "python tools/modbus_usb_client.py --port COM8 i2c-config 4 5",
            "python tools/modbus_usb_client.py --port COM8 i2c-scan",
            "python tools/modbus_usb_client.py --port COM8 spi-config 6 7 --miso 8 --cs 9",
            "python tools/modbus_usb_client.py --port COM8 bus-status",
        ],
    )
    add_body(doc, "pwm-set 参数依次为数字通道号、整数频率 Hz、占空比百分数，支持如 12.34 的占空比。默认启动；--disabled 保存配置但不启动，并停止当前通道已有 PWM。pwm-read 回读配置和使能，pwm-stop 停止并保留参数。")
    add_body(doc, "图形上位机 tools/modbus_usb_gui.py 包含数字 IO、ADC、PWM、I²C/SPI/UART 配置、串口助手和最多 4 路实时轮询波形。串口助手支持任意系统串口的文本或 HEX 收发、时间戳和日志保存；波形可导出 CSV。旧版固件仍可使用其已实现的功能。")

    build_heading = doc.add_heading("12 构建配置", level=1)
    build_heading.paragraph_format.keep_with_next = True
    build_intro = add_body(
        doc,
        "工程支持 ESP-IDF 5.4 和 5.5。首次构建时组件管理器下载 espressif/esp_tinyusb。菜单 USB Modbus IO"
        " 可设置从站地址、不完整帧超时和 GPIO35 到 GPIO37 开关。工程默认关闭控制台，避免 UART0 或"
        " USB 日志污染 Modbus 数据，并释放 GPIO43 和 GPIO44。默认启用两个 TinyUSB CDC 通道；"
        "本地已有旧 sdkconfig 时，应在 menuconfig 确认 CDC Channel Count 为 2。",
    )
    build_intro.paragraph_format.keep_together = True
    add_code(doc, ["idf.py set-target esp32s3", "idf.py menuconfig", "idf.py build", "idf.py -p COMx flash"])

    doc.add_heading("13 总线调试扩展", level=1)
    doc.add_heading("13 1 双虚拟串口和帧格式", level=2)
    add_body(doc, "CDC0 保留标准 Modbus RTU 从站、数字 IO、ADC 和 PWM。功能码 0x41 传输 I²C、SPI、UART 配置及 I²C/SPI 事务；不能向广播地址 0 发送。请求和成功响应均依次包含从站地址、0x41、操作码、数据长度、数据、CRC 低字节和 CRC 高字节。数据长度只计数据字段，频率和波特率为 32 位大端整数。最大 ADU 长度为 256 字节，错误响应使用标准 0xC1 异常帧。I²C 扫描可能耗时约 1 秒，主站响应超时建议至少 2 秒。")
    add_code(doc, ["请求/成功响应: SLAVE 41 OP LEN DATA... CRC_LO CRC_HI",
                   "异常响应:     SLAVE C1 EXCEPTION CRC_LO CRC_HI"])
    add_body(doc, "CDC1 是 UART1 原始字节透传口，不使用 Modbus 帧或 CRC。操作系统给两个 CDC 口分配的 COM 号不保证固定或相邻。电脑端对 CDC1 设置的波特率不自动改变 UART1，实际目标串口参数必须通过操作码 08 设置。")
    doc.add_heading("13 2 操作码", level=2)
    add_body(doc, "引脚字段均为第 5 节的数字通道号，不是 GPIO 编号。配置前引脚必须为模式 0 浮空输入；同一总线改配置前先释放。SPI 的 MISO 和 CS 可以用 FF 表示省略，同一配置中的其余引脚必须互不重复。外设占用的模式 6～8 不能通过普通模式寄存器写入或释放。")
    add_table(doc, ["操作码", "功能", "请求数据", "响应数据"], [
        ["01", "I²C 启用", "SDA SCL 频率4字节", "空"],
        ["02", "I²C 释放", "空", "空"],
        ["03", "I²C 扫描", "空", "16 字节地址位图"],
        ["04", "I²C 事务", "地址 写长度 读长度 写数据", "读数据"],
        ["05", "SPI 启用", "SCLK MOSI MISO CS 模式 频率4字节", "空"],
        ["06", "SPI 释放", "空", "空"],
        ["07", "SPI 全双工", "长度 发送数据", "同长度接收数据"],
        ["08", "UART 启用", "TX RX 波特率4字节 数据位 校验 停止位", "空"],
        ["09", "UART 释放", "空", "空"],
        ["0A", "状态查询", "空", "25 字节状态记录"],
    ], [0.75, 1.25, 3.55, 1.45])
    add_body(doc, "I²C 是 7 位地址主机，只接受 0x08～0x77。扫描只探测这 112 个地址；响应位图的 bitN 对应地址 N。事务写长、读长分别限 0～128，但不能同时为 0；同时写读时在两段之间使用重复起始条件，不发 STOP。单次事务超时 500 ms，扫描每个地址探测超时 5 ms。总线需要外部上拉，ESP32-S3 引脚不得承受 5 V。")
    add_body(doc, "SPI 使用 SPI2 主机、标准单线、MSB 优先和全双工模式。模式 0～3，频率 10 kHz～10 MHz；单次传输 1～128 字节，收发长度相同。配置 CS 时每次事务自动拉低并释放；省略 CS 时应自行控制目标设备片选。高频时必须在实际连线长度和目标器件上核验时序。")
    add_body(doc, "UART1 支持 300～2,000,000 bps，7 或 8 数据位，校验 0 无、1 偶、2 奇，停止位 1 或 2。没有硬件流控。CDC1 不打开时收到的 UART 字节会丢弃；队列和 UART 缓冲有限，持续高速发送或主机长时间不读可能丢字节，不能将它当作无损流记录器。")
    doc.add_heading("13 3 状态记录", level=2)
    add_body(doc, "操作码 0A 的响应数据固定 25 字节，偏移相对数据字段起算。未启用外设的引脚为 FF，频率为 0。")
    add_table(doc, ["偏移", "长度", "含义"], [
        ["0", "1", "标志位：bit0 I²C，bit1 SPI，bit2 UART"],
        ["1～2", "2", "I²C SDA、SCL 通道"],
        ["3～6", "4", "SPI SCLK、MOSI、MISO、CS 通道"],
        ["7～8", "2", "UART TX、RX 通道"],
        ["9～12", "4", "I²C 频率 Hz"],
        ["13～16", "4", "SPI 频率 Hz"],
        ["17～20", "4", "UART 波特率 bps"],
        ["21", "1", "SPI 模式"],
        ["22～24", "3", "UART 数据位、校验、停止位"],
    ], [0.95, 0.75, 5.3])
    doc.add_heading("13 4 报文示例", level=2)
    add_code(doc, [
        "I2C 100k:  01 41 01 06 04 05 00 01 86 A0 44 01",
        "I2C scan:  01 41 03 00 51 3C",
        "I2C W/R:   01 41 04 04 50 01 04 00 62 87",
        "SPI 1MHz:  01 41 05 09 06 07 08 09 00 00 0F 42 40 DE 39",
        "SPI xfer:  01 41 07 05 04 9F 00 00 00 6B DD",
        "UART 8N1:  01 41 08 09 19 1A 00 01 C2 00 08 00 01 19 38",
        "Status:    01 41 0A 00 57 6C",
    ])

    references_heading = doc.add_heading("14 参考资料", level=1)
    references_heading.paragraph_format.space_before = Pt(8)
    references_heading.paragraph_format.space_after = Pt(3)
    references = [
        ("ESP-IDF USB Device Stack", "https://docs.espressif.com/projects/esp-idf/en/v5.5/esp32s3/api-reference/peripherals/usb_device.html"),
        ("ESP-IDF ADC Oneshot Driver", "https://docs.espressif.com/projects/esp-idf/en/v5.5/esp32s3/api-reference/peripherals/adc_oneshot.html"),
        ("ESP-IDF LEDC Driver", "https://docs.espressif.com/projects/esp-idf/en/v5.5/esp32s3/api-reference/peripherals/ledc.html"),
        ("ESP-IDF I2C Master Driver", "https://docs.espressif.com/projects/esp-idf/en/v5.5/esp32s3/api-reference/peripherals/i2c.html"),
        ("ESP-IDF SPI Master Driver", "https://docs.espressif.com/projects/esp-idf/en/v5.5/esp32s3/api-reference/peripherals/spi_master.html"),
        ("ESP-IDF UART Driver", "https://docs.espressif.com/projects/esp-idf/en/v5.5/esp32s3/api-reference/peripherals/uart.html"),
        ("ESP32-S3 Datasheet", "https://documentation.espressif.com/esp32_s3_datasheet_en.pdf"),
        ("Modbus Application Protocol V1.1b3", "https://www.modbus.org/docs/Modbus_Application_Protocol_V1_1b3.pdf"),
        ("Modbus Serial Line Guide V1.02", "https://www.modbus.org/docs/Modbus_over_serial_line_V1_02.pdf"),
    ]
    for label, url in references:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        add_hyperlink(paragraph, label, url)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build the ESP32-S3 Modbus IO protocol DOCX")
    parser.add_argument("--board-image", type=Path, default=BOARD_IMAGE, help="Development board pinout image")
    args = parser.parse_args()
    build(args.board_image)
